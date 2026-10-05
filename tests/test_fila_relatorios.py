"""Os relatórios da fila (spec 2026-09-28-relatorios-da-fila-design).

O cliente: "um relatório que pode ser tirado pelo dono da conta, supervisor
ou gerente, e cada um é focado em uma coisa — o dono de tudo, o supervisor do
que ele é cadastrado, e o gerente no que ele é cadastrado —, diário, semanal,
mensal e mês passado". Todo número sai das contas do painel: se o papel e a
tela discordassem, ninguém saberia em qual acreditar.
"""

from datetime import datetime, timedelta
from decimal import Decimal

import pytest
from django.urls import reverse
from django.utils import timezone

from tests.fila_cenario import cadastros, logado, nova_loja, pessoa_na_loja, sylvia

pytestmark = pytest.mark.django_db


def local(*args):
    return timezone.make_aware(datetime(*args))


@pytest.fixture
def rede():
    from types import SimpleNamespace

    empresa, matriz, titular = sylvia()
    centro = nova_loja(empresa, "Centro")
    return SimpleNamespace(
        empresa=empresa, matriz=matriz, centro=centro, titular=titular,
        ana=pessoa_na_loja("ana", empresa, matriz),
        caio=pessoa_na_loja("caio", empresa, centro),
        gil=pessoa_na_loja("gil", empresa, centro, cargo="gerente"),
        sara=pessoa_na_loja("sara", empresa, None, cargo="supervisor"),
        cad=cadastros(empresa))


def atendimento(rede, pessoa, loja, fim, *, valor=None, motivo="padrao", midia=None,
                itens=()):
    """Um atendimento FECHADO em `fim`. `valor` None é não venda; `motivo=None`
    é a não venda sem motivo (a do ponto esquecido)."""
    from fila.models import Atendimento, ItemVendido, Presenca

    presenca = Presenca.irrestritos.create(empresa=rede.empresa, filial=loja,
                                          pessoa=pessoa, entrada=fim, saida=fim)
    a = Atendimento.irrestritos.create(
        empresa=rede.empresa, filial=loja, vendedor=pessoa, presenca=presenca,
        inicio=fim - timedelta(minutes=5), fim=fim,
        resultado="vendeu" if valor else "nao_vendeu",
        motivo=None if valor else (rede.cad.motivo if motivo == "padrao" else motivo),
        midia=midia, total=Decimal(valor or 0))
    for grupo, v in itens:
        ItemVendido.irrestritos.create(empresa=rede.empresa, atendimento=a,
                                       grupo=grupo, valor=Decimal(v))
    return a


def _recorte(rede, chave="mes", lojas=None, agora=None):
    from fila.indicadores import Recorte
    from fila.relatorio import periodo_escolhido

    return Recorte(rede.empresa, tuple(lojas or (rede.matriz, rede.centro)),
                   periodo_escolhido(chave, agora))


def _secao(relatorio, titulo):
    return next(s for s in relatorio.secoes if s.titulo == titulo)


class TestOPeriodo:
    def test_as_cinco_escolhas_e_o_padrao(self):
        from fila.relatorio import PADRAO, PERIODOS, periodo_escolhido

        assert [c for c, _r in PERIODOS] == ["hoje", "ontem", "7dias", "mes", "mes_passado"]
        assert PADRAO == "mes"
        assert periodo_escolhido("15dias").chave == "mes"
        assert periodo_escolhido(None).chave == "mes"


class TestOConteudo:
    def test_as_secoes_na_ordem(self, rede):
        from fila.relatorio import montar

        r = montar(_recorte(rede), rotulo_da_empresa=str(rede.empresa), gerado_por="Sylvia")
        assert [s.titulo for s in r.secoes] == [
            "Resumo", "Comparação", "Vendedores", "Motivos", "Mídias", "Grupos",
            "Pausas", "Lançamentos"]
        assert r.gerado_por == "Sylvia"
        assert "Matriz" in r.lojas and "Centro" in r.lojas

    def test_o_resumo_bate_com_o_painel(self, rede):
        from fila.indicadores import numeros
        from fila.relatorio import montar

        agora = timezone.now()
        atendimento(rede, rede.ana, rede.matriz, agora, valor="300")
        atendimento(rede, rede.caio, rede.centro, agora)
        recorte = _recorte(rede)
        r = montar(recorte, rotulo_da_empresa=str(rede.empresa), gerado_por="Sylvia")
        total = _secao(r, "Resumo").linhas[0]
        n = numeros(recorte)
        assert total.rotulo == "Total"
        assert (total.numeros.atendimentos, total.numeros.vendas, total.numeros.vendido) == (
            n.atendimentos, n.vendas, n.vendido)
        assert [l.rotulo for l in _secao(r, "Resumo").linhas[1:]] == ["Matriz", "Centro"]

    def test_vendedor_em_duas_lojas_tem_duas_linhas(self, rede):
        from fila.relatorio import montar
        from tests.conftest import alocar

        alocar(rede.ana, rede.empresa, "vendedor", filial=rede.centro)
        agora = timezone.now()
        atendimento(rede, rede.ana, rede.matriz, agora, valor="100")
        atendimento(rede, rede.ana, rede.centro, agora, valor="200")
        r = montar(_recorte(rede), rotulo_da_empresa="x", gerado_por="x")
        linhas = [(l["loja_nome"], l["nome"]) for l in _secao(r, "Vendedores").linhas]
        assert sorted(linhas) == [("Centro", "Ana"), ("Matriz", "Ana")]

    def test_o_pct_da_meta_so_nos_periodos_de_mes(self, rede):
        from fila.relatorio import montar

        agora = timezone.now()
        atendimento(rede, rede.ana, rede.matriz, agora, valor="100")
        for chave, tem in (("mes", True), ("7dias", False), ("hoje", False)):
            r = montar(_recorte(rede, chave), rotulo_da_empresa="x", gerado_por="x")
            rotulos = [c.rotulo for c in _secao(r, "Vendedores").colunas]
            assert ("% da meta" in rotulos) is tem, chave

    def test_a_pausa_da_gestao_fica_fora_da_pausa_do_vendedor(self, rede):
        from fila.models import Pausa, Presenca
        from fila.relatorio import montar

        agora = timezone.now().replace(microsecond=0)
        atendimento(rede, rede.ana, rede.matriz, agora, valor="100")
        presenca = Presenca.irrestritos.filter(pessoa=rede.ana).first()
        inicio = agora - timedelta(hours=2)
        Pausa.irrestritos.create(empresa=rede.empresa, pessoa=rede.ana, filial=rede.matriz,
                                 presenca=presenca, tipo=rede.cad.tipo, inicio=inicio,
                                 fim=inicio + timedelta(minutes=10))
        Pausa.irrestritos.create(empresa=rede.empresa, pessoa=rede.ana, filial=rede.matriz,
                                 presenca=presenca, fixa="gestao",
                                 inicio=inicio + timedelta(minutes=20),
                                 fim=inicio + timedelta(minutes=50))
        # "7dias", e não "hoje": as pausas começam duas horas atrás, e a suíte
        # rodando entre 00:00 e 02:00 as poria no dia anterior.
        r = montar(_recorte(rede, "7dias"), rotulo_da_empresa="x", gerado_por="x")
        ana = next(l for l in _secao(r, "Vendedores").linhas if l["nome"] == "Ana")
        assert ana["pausa"] == timedelta(minutes=10)
        pausas = dict(_secao(r, "Pausas").linhas)
        assert pausas == {"Gestão": 30, "Almoço": 10}

    def test_nao_venda_sem_motivo_no_lancamento(self, rede):
        from fila.relatorio import formatar, montar

        atendimento(rede, rede.ana, rede.matriz, timezone.now(), motivo=None)
        r = montar(_recorte(rede), rotulo_da_empresa="x", gerado_por="x")
        secao = _secao(r, "Lançamentos")
        coluna = next(c for c in secao.colunas if c.rotulo == "Motivo")
        assert formatar(coluna.valor(secao.linhas[0]), coluna.tipo) == "Fechado sem lançamento"

    def test_lancamentos_so_os_fechados_no_periodo(self, rede):
        from fila.indicadores import lancamentos_do_recorte
        from fila.models import Atendimento, Presenca

        agora = timezone.now()
        dentro = atendimento(rede, rede.ana, rede.matriz, agora, valor="100",
                             itens=[(rede.cad.grupo, "100")])
        atendimento(rede, rede.ana, rede.matriz, agora - timedelta(days=70), valor="50")
        presenca = Presenca.irrestritos.filter(pessoa=rede.caio).first() or \
            Presenca.irrestritos.create(empresa=rede.empresa, filial=rede.centro,
                                        pessoa=rede.caio, entrada=agora)
        Atendimento.irrestritos.create(empresa=rede.empresa, filial=rede.centro,
                                       vendedor=rede.caio, presenca=presenca, inicio=agora)
        assert [a.pk for a in lancamentos_do_recorte(_recorte(rede))] == [dentro.pk]

    def test_a_lista_de_lancamentos_nao_cresce_em_consultas(self, rede):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext

        from fila.relatorio import montar

        agora = timezone.now()

        def medir():
            with CaptureQueriesContext(connection) as c:
                r = montar(_recorte(rede), rotulo_da_empresa="x", gerado_por="x")
                secao = _secao(r, "Lançamentos")
                for linha in secao.linhas:
                    for coluna in secao.colunas:
                        coluna.valor(linha)
            return len(c)

        atendimento(rede, rede.ana, rede.matriz, agora, valor="100",
                    itens=[(rede.cad.grupo, "100")])
        poucas = medir()
        for _i in range(6):
            atendimento(rede, rede.caio, rede.centro, agora, valor="50",
                        itens=[(rede.cad.grupo2, "50")])
        assert medir() <= poucas

    def test_recorte_vazio_sai_com_as_secoes(self, rede):
        from fila.relatorio import montar

        r = montar(_recorte(rede, "hoje"), rotulo_da_empresa="x", gerado_por="x")
        assert len(r.secoes) == 8
        assert _secao(r, "Lançamentos").linhas == []
        assert _secao(r, "Resumo").linhas[0].numeros.atendimentos == 0

    def test_a_variacao_da_conversao_diz_pontos(self, rede):
        """Visto no PDF salvo pelo Chrome: a troca do ponto decimal pela
        vírgula pegava também os pontos da unidade, e a conversão saía
        "-7,5 p,p,"."""
        from fila.relatorio import montar

        # Datas fixas, com venda nos DOIS períodos: sem o anterior, a variação
        # é "—" e o teste não provaria nada.
        agora = local(2026, 9, 15, 12)
        atendimento(rede, rede.ana, rede.matriz, local(2026, 9, 10, 10), valor="100")
        atendimento(rede, rede.ana, rede.matriz, local(2026, 9, 10, 11))
        atendimento(rede, rede.ana, rede.matriz, local(2026, 8, 10, 10), valor="300")
        r = montar(_recorte(rede, "mes", agora=agora), rotulo_da_empresa="x",
                   gerado_por="x", agora=agora)
        linhas = {l[0]: l[3] for l in _secao(r, "Comparação").linhas}
        assert linhas["Conversão"] == "-50,0 p.p."
        assert linhas["Vendido"] == "-66,7 %"

    def test_a_variacao_escreve_a_virgula_no_numero(self):
        from fila.relatorio import _variacao_por_extenso

        assert _variacao_por_extenso(-7.5, "p.p.") == "-7,5 p.p."
        assert _variacao_por_extenso(12.25, "%") == "+12,2 %"

    def test_formatar(self):
        from fila.relatorio import formatar

        assert formatar(Decimal("1200"), "dinheiro") == "R$ 1.200,00"
        assert formatar(12.345, "porcento") == "12,3%"
        assert formatar(None, "porcento") == "—"
        assert formatar(timedelta(minutes=90), "minutos") == "90 min"
        assert formatar(3, "inteiro") == "3"


def test_os_cinco_periodos_em_castelhano():
    """"Ontem" e "Mês passado" ficaram de fora do `.po`, e o seletor saía com
    três opções em castelhano e duas em português (revisão final)."""
    from django.utils import translation

    from fila.relatorio import PERIODOS

    with translation.override("es"):
        rotulos = [str(r) for _chave, r in PERIODOS]
    assert rotulos == ["Diario (hoy)", "Ayer", "Semanal (últimos 7 días)",
                       "Mensual (este mes)", "Mes pasado"]


class TestOExcel:
    def _abrir(self, resposta):
        from io import BytesIO

        from openpyxl import load_workbook

        return load_workbook(BytesIO(resposta.content))

    def test_texto_digitado_nao_vira_formula(self, rede):
        """A observação da não venda é o vendedor quem digita, e o relatório
        vai para a mão do gerente. O openpyxl grava como FÓRMULA todo texto
        que começa com "=": um `=HYPERLINK(...)` na observação viraria um link
        vivo na planilha de quem a abre (revisão final de 28/09/2026). O
        nome da loja é dado do cliente também, e passa pela mesma porta."""
        from fila.relatorio import montar
        from fila.relatorio_saida import em_xlsx

        rede.centro.apelido = "=1+1"
        rede.centro.save()
        a = atendimento(rede, rede.caio, rede.centro, timezone.now())
        a.observacao = '=HYPERLINK("http://x";"ver")'
        a.save()
        livro = self._abrir(em_xlsx(montar(_recorte(rede), rotulo_da_empresa="x",
                                           gerado_por="x")))
        celulas = [c for folha in livro.worksheets for linha in folha.iter_rows()
                   for c in linha if isinstance(c.value, str) and c.value.startswith("=")]
        assert celulas, "o texto com = tem de chegar na planilha"
        assert {c.data_type for c in celulas} == {"s"}

    def test_as_abas_e_o_cabecalho(self, rede):
        from fila.relatorio import montar
        from fila.relatorio_saida import em_xlsx

        agora = timezone.now()
        atendimento(rede, rede.ana, rede.matriz, agora, valor="1200",
                    itens=[(rede.cad.grupo, "1200")])
        r = montar(_recorte(rede), rotulo_da_empresa="Sylvia Design", gerado_por="Sylvia")
        resposta = em_xlsx(r)
        assert resposta["Content-Type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        assert f'filename="{r.nome_do_arquivo}.xlsx"' in resposta["Content-Disposition"]
        livro = self._abrir(resposta)
        assert livro.sheetnames == ["Relatório", "Resumo", "Comparação", "Vendedores",
                                    "Motivos", "Mídias", "Grupos", "Pausas", "Lançamentos"]
        capa = {linha[0]: linha[1] for linha in livro["Relatório"].iter_rows(values_only=True)}
        assert capa["Empresa"] == "Sylvia Design" and capa["Gerado por"] == "Sylvia"

    def test_dinheiro_e_numero_e_nao_texto(self, rede):
        from fila.relatorio import montar
        from fila.relatorio_saida import em_xlsx

        atendimento(rede, rede.ana, rede.matriz, timezone.now(), valor="1200")
        livro = self._abrir(em_xlsx(montar(_recorte(rede), rotulo_da_empresa="x",
                                           gerado_por="x")))
        resumo = livro["Resumo"]
        cabecalho = [c.value for c in resumo[1]]
        vendido = resumo.cell(row=2, column=cabecalho.index("Vendido") + 1)
        conversao = resumo.cell(row=2, column=cabecalho.index("Conversão") + 1)
        assert vendido.value == 1200 and "R$" in vendido.number_format
        assert conversao.value == 1.0 and "%" in conversao.number_format

    def test_excel_do_recorte_vazio_abre(self, rede):
        from fila.relatorio import montar
        from fila.relatorio_saida import em_xlsx

        livro = self._abrir(em_xlsx(montar(_recorte(rede, "hoje"),
                                           rotulo_da_empresa="x", gerado_por="x")))
        assert livro["Lançamentos"].max_row == 1   # só o cabeçalho


class TestOPapel:
    def _setembro(self, rede):
        """Setembro até o dia 15 ao meio-dia, com venda nos DOIS períodos."""
        from fila.relatorio import montar

        agora = local(2026, 9, 15, 12)
        atendimento(rede, rede.ana, rede.matriz, local(2026, 9, 10, 10), valor="100")
        atendimento(rede, rede.caio, rede.centro, local(2026, 9, 10, 11))
        atendimento(rede, rede.ana, rede.matriz, local(2026, 8, 10, 10), valor="300")
        return montar(_recorte(rede, "mes", agora=agora), rotulo_da_empresa="Sylvia Design",
                      gerado_por="Sylvia", agora=agora)

    def test_a_chamada_do_mes(self, rede):
        """A capa abre com o NOME do período, e o intervalo e as lojas embaixo
        (28/09/2026: o cliente trocou a frase longa do resultado, que repetia
        os números dos cartões, por uma chamada)."""
        from fila.relatorio import chamada_do_periodo

        assert chamada_do_periodo(self._setembro(rede)) == (
            "Setembro de 2026", "De 1 a 15 de setembro de 2026 · Matriz e Centro")

    @pytest.mark.parametrize("chave, titulo, apoio", [
        ("hoje", "15 de setembro de 2026", "Matriz"),
        ("ontem", "14 de setembro de 2026", "Matriz"),
        ("7dias", "Últimos 7 dias", "De 9 a 15 de setembro de 2026 · Matriz"),
        ("mes_passado", "Agosto de 2026", "De 1 a 31 de agosto de 2026 · Matriz"),
    ])
    def test_a_chamada_de_cada_periodo(self, rede, chave, titulo, apoio):
        from fila.relatorio import chamada_do_periodo, montar

        agora = local(2026, 9, 15, 12)
        r = montar(_recorte(rede, chave, lojas=(rede.matriz,), agora=agora),
                   rotulo_da_empresa="x", gerado_por="x", agora=agora)
        assert chamada_do_periodo(r) == (titulo, apoio)

    def test_muitas_lojas_viram_numero(self):
        from fila.relatorio import _lojas_por_extenso

        assert _lojas_por_extenso(("A", "B", "C")) == "A, B e C"
        assert _lojas_por_extenso(("A", "B", "C", "D")) == "4 lojas"

    def test_o_papel(self, rede):
        from fila.relatorio_saida import em_impressao

        r = self._setembro(rede)
        html = em_impressao(r, com_logo=True).content.decode()
        assert '<h1 class="relatorio-chamada">Setembro de 2026</h1>' in html
        # O tipo do período, as datas e as lojas estão nos filtros dentro da
        # capa (a moldura padrão), e não mais numa pílula e numa linha de apoio.
        assert "<dd>Mensal (este mês)\n01/09/2026 a 15/09/2026</dd>" in html
        assert "<dd>Matriz, Centro</dd>" in html
        assert html.count('class="relatorio-numero"') == 5
        assert "▼" in html
        titulos = ["Vendas no período", "Lojas", "Vendedores", "Motivos de não venda", "Mídias",
                   "Grupos de item", "Pausas", "Lançamentos"]
        posicoes = [html.index(f"<h2>{t}</h2>") for t in titulos]
        assert posicoes == sorted(posicoes), "as seções na ordem, lançamentos por último"
        assert "Sylvia Design" in html and "Emitido em 15/09/2026 às 12:00 por Sylvia" in html
        assert "R$ 100,00" in html
        assert "window.print()" in html

    def test_o_grafico_por_dia(self, rede):
        """Uma coluna por dia do período, e o dia vazio aparece com zero: o
        gráfico que pula o dia ruim esconde justamente ele. A coluna mais alta
        é o maior dia, e só ela leva o destaque."""
        from fila.relatorio_saida import _papel

        dias = _papel(self._setembro(rede))["dias"]
        assert [d["rotulo"] for d in dias][:2] == ["01/09", "02/09"] and len(dias) == 15
        maior = [d for d in dias if d["maior"]]
        assert [d["rotulo"] for d in maior] == ["10/09"]
        assert 0 < maior[0]["altura"] <= 100
        assert all(d["altura"] == 0 for d in dias if d["rotulo"] != "10/09")

    def test_dia_sem_venda_nao_inventa_escala(self, rede):
        from fila.relatorio import montar
        from fila.relatorio_saida import _papel

        agora = local(2026, 9, 15, 12)
        r = montar(_recorte(rede, "hoje", agora=agora), rotulo_da_empresa="x",
                   gerado_por="x", agora=agora)
        p = _papel(r)
        assert p["marcas"] == [] and len(p["dias"]) == 24
        assert _papel(self._setembro(rede))["marcas"], "com venda, a escala aparece"

    def test_as_roscas(self, rede):
        """A conversão e a parte de cada loja no vendido saem como rosca; o
        pedaço de cada loja começa onde o da anterior terminou."""
        from fila.relatorio_saida import _papel

        p = _papel(self._setembro(rede))
        assert p["conversao"] == 50.0
        matriz, centro = p["lojas"]
        assert (matriz["fatia"], matriz["inicio"]) == (100.0, 0.0)
        assert (centro["fatia"], centro["inicio"]) == (0.0, 100.0)
        assert matriz["cor"] == "var(--chart-1)" and centro["cor"] == "var(--chart-2)"

    def test_o_podio(self, rede):
        """Os três primeiros no pódio, o resto na tabela a partir do 4º."""
        from fila.relatorio_saida import _papel, em_impressao

        r = self._setembro(rede)
        p = _papel(r)
        assert [v["posicao"] for v in p["podio"]] == [1, 2]
        assert p["podio"][0]["nome"] == "Ana" and p["resto"] == []
        html = em_impressao(r).content.decode()
        assert 'class="relatorio-podio"' in html
        assert 'class="relatorio-grafico"' in html
        assert html.count('class="relatorio-rosca"') == 2

    def test_sem_logo_nao_desenha_imagem_quebrada(self, rede):
        from fila.relatorio_saida import em_impressao

        r = self._setembro(rede)
        assert "/marca/empresa/menu" not in em_impressao(r, com_logo=False).content.decode()
        assert 'src="/marca/empresa/menu"' in em_impressao(r, com_logo=True).content.decode()

    def test_nome_com_html_sai_escapado(self, rede):
        from fila.relatorio import montar
        from fila.relatorio_saida import em_impressao

        rede.centro.apelido = "A&B <Centro>"
        rede.centro.save()
        atendimento(rede, rede.caio, rede.centro, timezone.now(), valor="10")
        html = em_impressao(montar(_recorte(rede), rotulo_da_empresa="x",
                                   gerado_por="x")).content.decode()
        assert "A&amp;B &lt;Centro&gt;" in html
        assert "<Centro>" not in html


class TestATela:
    def _get(self, cliente, **params):
        return cliente.get(reverse("fila_relatorios_gerais"), params)

    def _na_loja(self, cliente, loja):
        from plataforma.contexto import CHAVE

        sessao = cliente.session
        sessao[CHAVE] = loja.pk
        sessao.save()
        return cliente

    def test_quem_entra(self, rede):
        assert self._get(logado("sylvia")).status_code == 200
        assert self._get(logado("sara")).status_code == 200
        assert self._get(self._na_loja(logado("gil"), rede.centro)).status_code == 200
        assert self._get(logado("ana")).status_code == 404

    def test_os_gerais_sao_o_relatorio_visual(self, rede):
        """"Relatórios gerais" é o visual (28/09/2026): ele saiu da tela quando
        o de mídias entrou, e voltou com o menu de dois relatórios."""
        papel = self._get(logado("sylvia"), formato="impressao").content.decode()
        assert "<h2>Vendas no período</h2>" in papel
        planilha = self._get(logado("sylvia"), formato="xlsx")
        assert "relatorio-midias" not in planilha["Content-Disposition"]
        assert 'filename="relatorio-' in planilha["Content-Disposition"]

    def test_cada_tela_manda_para_ela_mesma(self, rede):
        html = self._get(logado("sylvia")).content.decode()
        assert 'action="/fila/relatorios/gerais"' in html and "Relatórios gerais" in html

    # O alcance pela tela e o Excel pela tela moram em
    # test_fila_relatorio_midias.py desde 28/09/2026: a tela passou a gerar o
    # relatório das lojas por mídia, e o visual ficou guardado (testado aqui
    # pelas funções, e não pela tela).

    def test_a_tela_sem_javascript(self, rede):
        html = self._get(logado("sylvia")).content.decode()
        assert '<form' in html and 'method="get"' in html
        for chave in ("hoje", "ontem", "7dias", "mes", "mes_passado"):
            assert f'value="{chave}"' in html
        # O formulário filtra; o PDF e o Excel ficam em cima do relatório
        # filtrado (05/10/2026, ver test_fila_relatorio_previa.py).
        assert 'name="ver" value="1"' in html
