"""O relatório das lojas por mídia (28/09/2026).

O cliente trouxe a planilha que a rede já usava — uma linha por loja com
atendimentos, vendas, quantos clientes vieram por cada canal e o
aproveitamento, e a linha da empresa embaixo — e pediu que a tela de
Relatórios gerasse ESTE, guardando o relatório visual para voltar depois.
"""

import pytest
from django.urls import reverse

from tests.fila_cenario import logado
from tests.test_fila_relatorios import _recorte, atendimento, local, rede  # noqa: F401

pytestmark = pytest.mark.django_db

AGORA = (2026, 9, 15, 12)


def _midia(rede, nome, ordem=0, ativo=True):
    from fila.models import Midia

    return Midia.irrestritos.create(empresa=rede.empresa, nome=nome, ordem=ordem, ativo=ativo)


def _montar(rede, chave="mes", lojas=None):
    from fila.relatorio_midias import montar

    agora = local(*AGORA)
    return montar(_recorte(rede, chave, lojas=lojas, agora=agora),
                  rotulo_da_empresa="Sylvia Design", gerado_por="Sylvia", agora=agora)


@pytest.fixture
def canais(rede):
    """Fachada e Instagram ativas, TV desativada; no Centro, um cliente de
    Instagram que comprou e um sem mídia; na Matriz, dois de Fachada (um
    comprou) e um da TV, que foi desativada depois."""
    fachada = _midia(rede, "Fachada", ordem=1)
    instagram = _midia(rede, "Instagram", ordem=2)
    tv = _midia(rede, "TV", ordem=3, ativo=False)
    _midia(rede, "Rádio", ordem=4, ativo=False)
    dia = local(2026, 9, 10, 10)
    atendimento(rede, rede.ana, rede.matriz, dia, valor="100", midia=fachada)
    atendimento(rede, rede.ana, rede.matriz, dia, midia=fachada)
    atendimento(rede, rede.ana, rede.matriz, dia, midia=tv)
    atendimento(rede, rede.caio, rede.centro, dia, valor="50", midia=instagram)
    atendimento(rede, rede.caio, rede.centro, dia)
    return rede


class TestAsContas:
    def test_as_colunas_de_midia(self, canais):
        """As mídias ATIVAS na ordem do cadastro; a desativada só entra se
        trouxe alguém no período (a TV sim, o Rádio não); e "Sem mídia" no
        fim, porque houve atendimento sem ela."""
        r = _montar(canais)
        assert r.midias == ("Fachada", "Instagram", "TV", "Sem mídia")

    def test_uma_linha_por_loja(self, canais):
        r = _montar(canais)
        matriz, centro = r.linhas
        assert (matriz.nome, matriz.atendimentos, matriz.vendas, matriz.por_midia) == (
            "Matriz", 3, 1, (2, 0, 1, 0))
        assert (centro.nome, centro.atendimentos, centro.vendas, centro.por_midia) == (
            "Centro", 2, 1, (0, 1, 0, 1))
        assert round(matriz.aproveitamento, 2) == 33.33
        assert centro.aproveitamento == 50.0

    def test_a_linha_da_empresa(self, canais):
        total = _montar(canais).total
        assert (total.atendimentos, total.vendas, total.por_midia) == (5, 2, (2, 1, 1, 1))
        assert total.aproveitamento == 40.0

    def test_as_midias_somam_os_atendimentos(self, canais):
        """Com "Sem mídia", as colunas de canal somam os atendimentos de cada
        loja — sem ela, a planilha antiga não fechava (52 de 66 na Zakinha)."""
        for linha in (*_montar(canais).linhas, _montar(canais).total):
            assert sum(linha.por_midia) == linha.atendimentos

    def test_bate_com_o_painel(self, canais):
        from fila import indicadores as ind

        recorte = _recorte(canais, agora=local(*AGORA))
        n = ind.numeros(recorte)
        total = _montar(canais).total
        assert (total.atendimentos, total.vendas) == (n.atendimentos, n.vendas)
        painel = {nome: at for nome, at, _vendas in ind.midias(recorte)}
        r = _montar(canais)
        assert dict(zip(r.midias, total.por_midia)) == painel

    def test_loja_sem_atendimento_aparece_zerada(self, rede):
        _midia(rede, "Fachada")
        r = _montar(rede)
        assert [(l.nome, l.atendimentos, l.por_midia, l.aproveitamento) for l in r.linhas] == [
            ("Matriz", 0, (0,), None), ("Centro", 0, (0,), None)]
        assert r.midias == ("Fachada",), "sem atendimento sem mídia, não há a coluna"

    def test_aberto_e_fora_do_periodo_nao_contam(self, canais):
        """A regra do painel: entra pela hora do FIM, e o do mês passado não."""
        atendimento(canais, canais.ana, canais.matriz, local(2026, 8, 31, 10))
        assert _montar(canais).total.atendimentos == 5


class TestAsSaidas:
    def test_o_excel(self, canais):
        from io import BytesIO

        from openpyxl import load_workbook

        from fila.relatorio_midias import em_xlsx

        resposta = em_xlsx(_montar(canais))
        assert "relatorio-midias-sylvia-design-mes-" in resposta["Content-Disposition"]
        folha = load_workbook(BytesIO(resposta.content)).active
        linhas = [[c.value for c in linha] for linha in folha.iter_rows()]
        cabecalho = next(l for l in linhas if l and l[0] == "Loja")
        assert cabecalho == ["Loja", "Atendimentos", "Vendas", "Fachada", "Instagram", "TV",
                             "Sem mídia", "Aproveitamento"]
        matriz = next(l for l in linhas if l and l[0] == "Matriz")
        assert matriz[:7] == ["Matriz", 3, 1, 2, 0, 1, 0]
        assert abs(matriz[7] - 1 / 3) < 1e-9, "aproveitamento é número, em fração"
        empresa = linhas[-1]
        assert empresa[0] == "Aproveitamento Empresa" and empresa[1:3] == [5, 2]
        assert empresa[7] == 0.4

    def test_o_papel(self, canais):
        from fila.relatorio_midias import em_impressao

        html = em_impressao(_montar(canais)).content.decode()
        assert 'class="rm-midia">Fachada<span>40,0%</span>' in html
        assert 'class="rm-midia">Sem mídia<span>' in html
        assert "Aproveitamento Empresa" in html
        assert "33,33%" in html and "40,00%" in html
        assert "size: A4 landscape" in html or "relatorio-midias" in html
        assert "window.print()" in html

    def test_o_mapa_de_calor(self, canais):
        """O tom de cada célula é a parte da mídia nos atendimentos DA LOJA
        (28/09/2026, "dá para deixar mais bonitinho?"): a Matriz teve 2 de 3
        pela Fachada, e essa célula é a mais forte da linha; zero não tem tom."""
        from fila.relatorio_midias import _papel

        p = _papel(_montar(canais))
        matriz = p["linhas"][0]
        fachada, instagram = matriz["celulas"][0], matriz["celulas"][1]
        assert (fachada["n"], round(fachada["parte"], 1)) == (2, 66.7)
        assert (instagram["n"], instagram["parte"], instagram["tom"]) == (0, 0.0, 0)
        assert fachada["tom"] > matriz["celulas"][2]["tom"] > 0

    def test_a_parte_de_cada_midia_na_rede(self, canais):
        from fila.relatorio_midias import _papel

        cabecalho = _papel(_montar(canais))["midias"]
        assert [(c["nome"], c["parte"]) for c in cabecalho] == [
            ("Fachada", "40,0%"), ("Instagram", "20,0%"), ("TV", "20,0%"), ("Sem mídia", "20,0%")]

    def test_quem_passa_o_aproveitamento_da_empresa(self, canais):
        """O Centro (50%) passa a empresa (40%), a Matriz (33%) não; a barra é
        relativa à loja que mais aproveitou."""
        from fila.relatorio_midias import _papel

        matriz, centro = _papel(_montar(canais))["linhas"]
        assert (matriz["acima"], centro["acima"]) == (False, True)
        assert centro["barra"] == 100.0 and round(matriz["barra"], 1) == 66.7

    def test_nome_com_html_sai_escapado(self, rede):
        from fila.relatorio_midias import em_impressao

        _midia(rede, "<b>Rua</b>")
        html = em_impressao(_montar(rede)).content.decode()
        assert "&lt;b&gt;Rua&lt;/b&gt;" in html and "<b>Rua</b>" not in html


class TestATela:
    """A tela de Relatórios gera ESTE relatório; o visual ficou guardado."""

    def _get(self, cliente, **params):
        return cliente.get(reverse("fila_relatorios"), params)

    def _na_loja(self, cliente, loja):
        from plataforma.contexto import CHAVE

        sessao = cliente.session
        sessao[CHAVE] = loja.pk
        sessao.save()
        return cliente

    def test_o_pdf_da_tela_e_o_das_midias(self, rede):
        papel = self._get(logado("sylvia"), formato="impressao").content.decode()
        assert "Aproveitamento Empresa" in papel

    def test_o_excel_da_tela_e_o_das_midias(self, rede):
        resposta = self._get(logado("sylvia"), formato="xlsx", periodo="mes_passado")
        assert "relatorio-midias-" in resposta["Content-Disposition"]
        assert "mes-passado" in resposta["Content-Disposition"]

    def test_o_gerente_so_ve_a_loja_dele(self, rede):
        """A loja forjada cai na dele: a tabela sai só com o Centro."""
        gil = self._na_loja(logado("gil"), rede.centro)
        papel = self._get(gil, formato="impressao", loja=str(rede.matriz.pk)).content.decode()
        assert "Centro" in papel and "Matriz" not in papel

    def test_o_dono_ve_todas_e_escolhe_uma(self, rede):
        dono = logado("sylvia")
        papel = self._get(dono, formato="impressao").content.decode()
        assert "Matriz" in papel and "Centro" in papel
        so_centro = self._get(dono, formato="impressao", loja=str(rede.centro.pk)).content.decode()
        assert "Centro" in so_centro and "Matriz" not in so_centro

    def test_o_supervisor_ve_as_lojas_da_empresa(self, rede):
        papel = self._get(logado("sara"), formato="impressao").content.decode()
        assert "Matriz" in papel and "Centro" in papel
