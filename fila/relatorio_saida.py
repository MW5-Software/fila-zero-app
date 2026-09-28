"""O relatório no papel e na planilha (28/09/2026).

Os dois leem o MESMO `fila.relatorio.Relatorio`. O Excel grava NÚMERO nas
colunas de número (somável, filtrável), com o formato de moeda e de
porcentagem na célula; o papel escreve o texto com `relatorio.formatar`.
"""

from __future__ import annotations

from datetime import timedelta
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone
from django.utils.translation import gettext as _

from comum.exportacao import CONTENT_TYPE_XLSX

from .relatorio import Relatorio, formatar

__all__ = ["em_impressao", "em_xlsx"]

_FORMATO = {"dinheiro": 'R$ #,##0.00', "porcento": "0.0%", "inteiro": "0",
            "minutos": "0"}


def _celula(valor, tipo):
    """O valor que a célula GRAVA. Porcentagem vai como fração (12,5% é
    0,125), que é o que o formato `0.0%` do Excel espera."""
    if valor is None or valor == "":
        return None
    if tipo == "dinheiro":
        return float(Decimal(valor))
    if tipo == "porcento":
        return float(valor) / 100
    if tipo == "minutos":
        return int(valor.total_seconds() // 60) if isinstance(valor, timedelta) else int(valor)
    if tipo == "inteiro":
        return int(valor)
    return str(valor)


def em_xlsx(relatorio: Relatorio) -> HttpResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    livro = Workbook()
    capa = livro.active
    capa.title = _("Relatório")
    for rotulo, valor in (
            (_("Relatório"), relatorio.titulo), (_("Empresa"), relatorio.empresa),
            (_("Lojas"), relatorio.lojas), (_("Período"), relatorio.periodo),
            (_("Gerado por"), relatorio.gerado_por),
            (_("Gerado em"), timezone.localtime(relatorio.gerado_em).strftime("%d/%m/%Y %H:%M"))):
        capa.append([rotulo, valor])
    for secao in relatorio.secoes:
        # O formato limita o nome da aba a 31 caracteres.
        folha = livro.create_sheet(secao.titulo[:31])
        folha.append([c.rotulo for c in secao.colunas])
        for celula in folha[1]:
            celula.font = Font(bold=True)
        for linha in secao.linhas:
            folha.append([_celula(c.valor(linha), c.tipo) for c in secao.colunas])
        for indice, coluna in enumerate(secao.colunas, start=1):
            formato = _FORMATO.get(coluna.tipo)
            if formato:
                for (celula,) in folha.iter_rows(min_row=2, min_col=indice, max_col=indice):
                    celula.number_format = formato
    resposta = HttpResponse(content_type=CONTENT_TYPE_XLSX)
    resposta["Content-Disposition"] = f'attachment; filename="{relatorio.nome_do_arquivo}.xlsx"'
    livro.save(resposta)
    return resposta
