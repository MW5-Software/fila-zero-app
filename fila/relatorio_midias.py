"""O relatório das lojas por mídia (28/09/2026).

O cliente trouxe a planilha que a rede já montava à mão: uma linha por loja
com atendimentos, vendas, quantos clientes vieram por cada canal ("Clientes
Fachada", "Clientes Instagram", "Passagem"…) e o aproveitamento, e a linha
da empresa embaixo. A tela de Relatórios passou a gerar ESTE; o relatório
visual (`fila/relatorio.py`, `fila/relatorio_saida.py`) continua no código e
nos testes, guardado para voltar à tela depois.

Os números saem da mesma regra do painel (`indicadores._atendimentos`: entra
pela hora do FIM, aberto não entra), e as colunas de mídia somam os
atendimentos — por isso "Sem mídia" aparece quando há atendimento sem ela:
sem a coluna, a planilha antiga não fechava (52 de 66 numa loja).
"""

from dataclasses import dataclass
from datetime import datetime

from django.db.models import Count, Q
from django.http import HttpResponse
from django.utils import timezone
from django.utils.text import slugify
from django.utils.translation import gettext as _

from . import indicadores as ind
from .models import Midia, Resultado
from .relatorio import _periodo_por_extenso

__all__ = ["Linha", "RelatorioDeMidias", "em_impressao", "em_xlsx", "montar"]


@dataclass(frozen=True)
class Linha:
    nome: str
    atendimentos: int
    vendas: int
    #: Atendimentos por mídia, na ordem de `RelatorioDeMidias.midias`.
    por_midia: tuple
    #: Vendas ÷ atendimentos, em porcentagem; None sem atendimento, como a
    #: conversão do painel ("—", e não 0%).
    aproveitamento: "float | None"


@dataclass(frozen=True)
class RelatorioDeMidias:
    titulo: str
    empresa: str
    lojas: str
    periodo: str
    gerado_por: str
    gerado_em: datetime
    midias: tuple
    linhas: tuple
    total: Linha
    nome_do_arquivo: str


def _aproveitamento(vendas, atendimentos):
    return vendas * 100 / atendimentos if atendimentos else None


def _colunas(recorte, usadas: set) -> tuple:
    """As mídias ATIVAS da empresa na ordem do cadastro, as desativadas que
    trouxeram alguém no período (o canal saiu do cadastro, e o cliente que
    veio por ele continua contado) e, no fim, "Sem mídia" se houve
    atendimento sem ela. Por NOME: em "Todas as empresas" o "Instagram" de
    cada empresa é o mesmo canal, e duas colunas iguais não diriam nada."""
    nomes = []
    for m in ind.do_recorte(Midia, recorte).order_by("ordem", "nome"):
        if (m.ativo or m.nome in usadas) and m.nome not in nomes:
            nomes.append(m.nome)
    if None in usadas:
        nomes.append(str(ind.SEM_MIDIA))
    return tuple(nomes)


def montar(recorte, *, rotulo_da_empresa: str, gerado_por: str, agora=None) -> RelatorioDeMidias:
    agora = agora or timezone.now()
    contagem = {}
    for linha in (ind._atendimentos(recorte)
                  .values("filial_id", "midia__nome")
                  .annotate(n=Count("pk"),
                            vendas=Count("pk", filter=Q(resultado=Resultado.VENDEU)))):
        contagem[(linha["filial_id"], linha["midia__nome"])] = (linha["n"], linha["vendas"])
    midias = _colunas(recorte, {nome for _loja, nome in contagem})
    chaves = [None if m == str(ind.SEM_MIDIA) else m for m in midias]

    def linha_de(nome, lojas):
        daqui = {chave: nv for chave, nv in contagem.items() if chave[0] in lojas}
        atendimentos = sum(n for n, _v in daqui.values())
        vendas = sum(v for _n, v in daqui.values())
        por_midia = tuple(sum(n for (_l, m), (n, _v) in daqui.items() if m == chave)
                          for chave in chaves)
        return Linha(nome, atendimentos, vendas, por_midia, _aproveitamento(vendas, atendimentos))

    linhas = tuple(linha_de(str(l), [l.pk]) for l in recorte.lojas)
    total = linha_de(_("Aproveitamento Empresa"), [l.pk for l in recorte.lojas])
    return RelatorioDeMidias(
        # O nome do item do menu (28/09/2026, pedido do cliente); era "Lojas
        # por mídia".
        titulo=_("Relatórios por mídia"), empresa=rotulo_da_empresa,
        lojas=", ".join(str(l) for l in recorte.lojas),
        periodo=_periodo_por_extenso(recorte.periodo),
        gerado_por=gerado_por, gerado_em=agora, midias=midias, linhas=linhas, total=total,
        nome_do_arquivo=slugify(f"relatorio-midias-{rotulo_da_empresa}-"
                                f"{recorte.periodo.chave.replace('_', '-')}-"
                                f"{timezone.localtime(agora):%Y-%m-%d}"))


def porcento(valor) -> str:
    """Duas casas, como a planilha que a rede usava ("16,67%")."""
    return "—" if valor is None else f"{valor:.2f}%".replace(".", ",")


def _gerado_em(r) -> str:
    quando = timezone.localtime(r.gerado_em)
    return _("%(data)s às %(hora)s") % {"data": quando.strftime("%d/%m/%Y"),
                                        "hora": quando.strftime("%H:%M")}


def em_xlsx(r: RelatorioDeMidias) -> HttpResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    from comum.exportacao import CONTENT_TYPE_XLSX

    from .relatorio_saida import _nada_e_formula

    livro = Workbook()
    folha = livro.active
    folha.title = r.titulo[:31]
    for rotulo, valor in ((_("Relatório"), r.titulo), (_("Empresa"), r.empresa),
                          (_("Lojas"), r.lojas), (_("Período"), r.periodo),
                          (_("Gerado por"), r.gerado_por), (_("Gerado em"), _gerado_em(r))):
        folha.append([rotulo, valor])
    folha.append([])
    folha.append([_("Loja"), _("Atendimentos"), _("Vendas"), *r.midias, _("Aproveitamento")])
    for celula in folha[folha.max_row]:
        celula.font = Font(bold=True)
    for linha in (*r.linhas, r.total):
        # O aproveitamento vai como FRAÇÃO com formato de porcentagem: número
        # de verdade, que quem abre a planilha soma e compara.
        folha.append([linha.nome, linha.atendimentos, linha.vendas, *linha.por_midia,
                      None if linha.aproveitamento is None else linha.aproveitamento / 100])
        folha.cell(folha.max_row, folha.max_column).number_format = "0.00%"
    for celula in folha[folha.max_row]:
        celula.font = Font(bold=True)
    _nada_e_formula(livro)
    resposta = HttpResponse(content_type=CONTENT_TYPE_XLSX)
    resposta["Content-Disposition"] = f'attachment; filename="{r.nome_do_arquivo}.xlsx"'
    livro.save(resposta)
    return resposta


def _papel(r: RelatorioDeMidias) -> dict:
    """O que o papel desenha além dos números (28/09/2026, "dá para deixar
    mais bonitinho?"): a tabela vira mapa de calor. O tom de cada célula é a
    parte daquela mídia nos atendimentos DA LOJA — a pergunta do gerente é
    "de onde vem o cliente desta loja", e o tom responde sem ler número. O
    cabeçalho de cada mídia diz a parte dela na rede, e o aproveitamento
    ganha a barra, relativa à loja que mais aproveitou."""
    maior = max((l.aproveitamento or 0 for l in r.linhas), default=0) or 1
    empresa = r.total.aproveitamento

    def celula(n, atendimentos):
        parte = n * 100 / atendimentos if atendimentos else 0.0
        # A cor da marca na proporção da parte, com teto de 85%: acima disso
        # o fundo engole o número. Com a parte crua (0,7 × parte na primeira
        # versão) o canal de metade da loja saía num azul pálido e o mapa não
        # se lia; a partir de 45% a letra passa a clara (`.rm-forte`).
        tom = min(85, round(parte * 1.15))
        return {"n": n, "parte": parte, "tom": tom}

    return {
        "midias": [{"nome": m, "parte": porcento_curto(
            n * 100 / r.total.atendimentos if r.total.atendimentos else None)}
            for m, n in zip(r.midias, r.total.por_midia)],
        "linhas": [{"linha": l,
                    "celulas": [celula(n, l.atendimentos) for n in l.por_midia],
                    "barra": (l.aproveitamento or 0) * 100 / maior,
                    "acima": (l.aproveitamento is not None and empresa is not None
                              and l.aproveitamento > empresa)}
                   for l in r.linhas],
    }


def porcento_curto(valor) -> str:
    return "—" if valor is None else f"{valor:.1f}%".replace(".", ",")


def em_impressao(r: RelatorioDeMidias, *, com_logo: bool = False,
                 lojas: "str | None" = None) -> HttpResponse:
    """A moldura é a mesma do relatório geral (`fila/moldura.py`); `com_logo`
    e `lojas` têm o sentido de lá."""
    from .moldura import moldura_de

    from django.utils.translation import get_language

    from comum.estaticos import versionado

    from .ambiente import ambiente_da_fila

    html = ambiente_da_fila().get_template("fila/relatorio_midias.html").render(
        r=r, p=_papel(r), porcento=porcento, gerado_em=_gerado_em(r), idioma=get_language() or "pt-BR",
        mo=moldura_de(empresa=r.empresa, periodo=r.periodo, lojas=lojas or r.lojas,
                      gerado_por=r.gerado_por, gerado_em=r.gerado_em, com_logo=com_logo),
        folha_da_moldura=versionado("/static/fila/moldura.css"),
        folha_do_sistema=versionado("/static/nucleo/mw5.css"),
        folha_da_impressao=versionado("/static/plataforma/impressao.css"),
        folha_do_relatorio=versionado("/static/fila/relatorio_midias.css"))
    return HttpResponse(html)
